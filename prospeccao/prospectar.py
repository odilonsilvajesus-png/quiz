#!/usr/bin/env python3
"""Prospecção na base de seguidores: quem entre os seus seguidores pode virar cliente.

Etapas (rode nesta ordem; cada uma grava em prospeccao/dados/<conta>/):
    1. seguidores   lê a exportação oficial do Instagram (sua conta) → seguidores.csv
    2. engajados    (opcional, Apify) quem comentou nos seus últimos posts → engajados.csv
    3. enriquecer   (Apify) dados públicos de cada perfil: bio, categoria, link, porte → perfis.jsonl
    4. pontuar      aplica o ICP (icp.json) → leads.csv, leads.xlsx, resumo.md

Exemplos:
    python prospeccao/prospectar.py seguidores ~/Downloads/instagram-odilon.mentor.zip
    python prospeccao/prospectar.py engajados --posts 30 --confirmar
    python prospeccao/prospectar.py enriquecer --limite 300              # só mostra quanto vai coletar
    python prospeccao/prospectar.py enriquecer --limite 300 --confirmar  # coleta (gasta créditos Apify)
    python prospeccao/prospectar.py pontuar

A coleta usa só dados PÚBLICOS de perfil. Exige a variável de ambiente APIFY_TOKEN (nunca em arquivo).
"""

import argparse
import csv
import io
import json
import os
import re
import sys
import unicodedata
import urllib.parse
import urllib.error
import urllib.request
import zipfile
from datetime import datetime, timezone
from pathlib import Path

AQUI = Path(__file__).resolve().parent
ICP_PADRAO = AQUI / "icp.json"
APIFY = "https://api.apify.com/v2/acts/apify~instagram-scraper/run-sync-get-dataset-items"
RESERVADOS = {"", "p", "reel", "reels", "stories", "explore", "accounts", "direct", "_u", "tv"}


# ---------- utilidades ----------
def pasta_conta(conta: str) -> Path:
    p = AQUI / "dados" / conta.lstrip("@").lower()
    p.mkdir(parents=True, exist_ok=True)
    return p


def norm(txt) -> str:
    txt = unicodedata.normalize("NFKD", str(txt or "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", txt.lower()).strip()


def achar(texto: str, palavras: list[str]) -> list[str]:
    """Palavras (sem acento, minúsculas) encontradas como palavra inteira; 're:' = expressão regular."""
    achadas = []
    for p in palavras:
        padrao = p[3:] if p.startswith("re:") else r"(?<![a-z0-9])" + re.escape(norm(p)) + r"(?![a-z0-9])"
        if re.search(padrao, texto):
            achadas.append(p[3:] if p.startswith("re:") else p)
    return achadas


def usuario_de(valor: str) -> str:
    v = (valor or "").strip()
    if "instagram.com" in v:
        partes = [x for x in urllib.parse.urlparse(v if "://" in v else "https://" + v).path.split("/") if x]
        partes = [x for x in partes if x.lower() not in RESERVADOS]
        v = partes[0] if partes else ""
    return v.lstrip("@").strip().lower()


def ler_csv(caminho: Path) -> list[dict]:
    if not caminho.exists():
        return []
    with caminho.open(encoding="utf-8") as f:
        return list(csv.DictReader(f))


def gravar_csv(caminho: Path, linhas: list[dict], campos: list[str]) -> None:
    with caminho.open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=campos, extrasaction="ignore")
        w.writeheader()
        w.writerows(linhas)


def token() -> str:
    t = os.environ.get("APIFY_TOKEN")
    if not t:
        sys.exit("Falta a variável APIFY_TOKEN (export APIFY_TOKEN=...). Não coloque o token em arquivo.")
    return t


def apify(payload: dict) -> list[dict]:
    url = f"{APIFY}?token={urllib.parse.quote(token())}&timeout=900"
    req = urllib.request.Request(url, json.dumps(payload).encode(), {"Content-Type": "application/json"})
    with urllib.request.urlopen(req, timeout=960) as r:
        return json.loads(r.read().decode())


# ---------- 1. seguidores ----------
def _seguidores_json(dados) -> list[tuple[str, int | None]]:
    if isinstance(dados, dict):  # algumas versões embrulham numa chave
        dados = next((v for v in dados.values() if isinstance(v, list)), [])
    saida = []
    for item in dados:
        sld = (item.get("string_list_data") or [{}])[0]
        u = sld.get("value") or item.get("title") or usuario_de(sld.get("href", ""))
        saida.append((usuario_de(u), sld.get("timestamp")))
    return saida


MESES = {m: i for i, m in enumerate(["jan", "fev", "mar", "abr", "mai", "jun", "jul", "ago", "set", "out", "nov", "dez"], 1)}
MESES.update({m: i for i, m in enumerate(["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)})


def _data_html(txt: str) -> int | None:
    m = re.search(r"([a-z]{3})\w*\.? (\d{1,2}), (\d{4})", norm(txt))
    if m and m.group(1) in MESES:
        return int(datetime(int(m.group(3)), MESES[m.group(1)], int(m.group(2)), tzinfo=timezone.utc).timestamp())
    return None


def _seguidores_html(html: str) -> list[tuple[str, int | None]]:
    saida = []
    blocos = re.split(r'(?=<a [^>]*href="https://www\.instagram\.com/)', html)
    for b in blocos[1:]:
        h = re.match(r'<a [^>]*href="([^"]+)"', b).group(1)
        resto = re.sub(r"<[^>]+>", " ", b.split("</a>", 1)[-1])
        saida.append((usuario_de(h), _data_html(resto)))
    return saida


def cmd_seguidores(a) -> None:
    origem = Path(a.arquivo).expanduser()
    brutos: list[tuple[str, int | None]] = []

    def processar(nome: str, conteudo: str) -> None:
        base = nome.rsplit("/", 1)[-1].lower()
        if not base.startswith("followers") and not base.startswith("seguidores"):
            return
        if base.endswith(".json"):
            brutos.extend(_seguidores_json(json.loads(conteudo)))
        elif base.endswith(".html"):
            brutos.extend(_seguidores_html(conteudo))
        elif base.endswith(".csv"):
            for r in csv.DictReader(io.StringIO(conteudo)):
                u = r.get("usuario") or r.get("username") or r.get("Username") or next(iter(r.values()), "")
                brutos.append((usuario_de(u), None))

    if origem.suffix.lower() == ".zip":
        with zipfile.ZipFile(origem) as z:
            for n in z.namelist():
                processar(n, z.read(n).decode("utf-8", "ignore"))
    elif origem.is_dir():
        for f in origem.rglob("*"):
            if f.is_file():
                processar(f.name, f.read_text(encoding="utf-8", errors="ignore"))
    else:
        nome = origem.name if origem.name.lower().startswith(("followers", "seguidores")) else "followers_" + origem.name
        processar(nome, origem.read_text(encoding="utf-8", errors="ignore"))

    vistos: dict[str, int | None] = {}
    for u, ts in brutos:
        if u and u not in vistos:
            vistos[u] = ts
    if not vistos:
        sys.exit("Nenhum seguidor encontrado. Confira se a exportação tem 'Seguidores e seguindo' (arquivo followers_1.json).")
    linhas = [{"usuario": u, "seguiu_em": datetime.fromtimestamp(ts, timezone.utc).date().isoformat() if ts else ""}
              for u, ts in sorted(vistos.items(), key=lambda kv: -(kv[1] or 0))]
    destino = pasta_conta(a.conta) / "seguidores.csv"
    gravar_csv(destino, linhas, ["usuario", "seguiu_em"])
    print(f"{len(linhas)} seguidores → {destino}")


# ---------- 2. engajados ----------
def cmd_engajados(a) -> None:
    conta = a.conta.lstrip("@").lower()
    if not a.confirmar:
        print(f"Vai coletar os comentários dos últimos {a.posts} posts de @{conta} (até {a.por_post} por post) no Apify.\n"
              f"Isso gasta créditos. Rode de novo com --confirmar.")
        return
    posts = apify({"directUrls": [f"https://www.instagram.com/{conta}/"], "resultsType": "posts", "resultsLimit": a.posts})
    urls = [p["url"] for p in posts if p.get("url")]
    print(f"{len(urls)} posts; coletando comentários…", file=sys.stderr)
    coment = apify({"directUrls": urls, "resultsType": "comments", "resultsLimit": a.por_post}) if urls else []
    contagem: dict[str, dict] = {}
    for c in coment:
        u = (c.get("ownerUsername") or "").lower()
        if not u or u == conta:
            continue
        d = contagem.setdefault(u, {"usuario": u, "comentarios": 0, "ultimo": ""})
        d["comentarios"] += 1
        d["ultimo"] = max(d["ultimo"], (c.get("timestamp") or "")[:10])
    linhas = sorted(contagem.values(), key=lambda d: -d["comentarios"])
    destino = pasta_conta(conta) / "engajados.csv"
    gravar_csv(destino, linhas, ["usuario", "comentarios", "ultimo"])
    print(f"{len(linhas)} perfis que comentaram → {destino}")


# ---------- 3. enriquecer ----------
def carregar_perfis(pasta: Path) -> dict[str, dict]:
    arq = pasta / "perfis.jsonl"
    perfis = {}
    if arq.exists():
        for linha in arq.read_text(encoding="utf-8").splitlines():
            if linha.strip():
                d = json.loads(linha)
                perfis[d["usuario"]] = d
    return perfis


def cmd_enriquecer(a) -> None:
    pasta = pasta_conta(a.conta)
    seguidores = ler_csv(pasta / "seguidores.csv")
    engajados = ler_csv(pasta / "engajados.csv")
    if not seguidores and not engajados:
        sys.exit("Rode antes a etapa 'seguidores' (e/ou 'engajados').")
    ja = carregar_perfis(pasta)
    # prioridade: quem comentou → seguidores mais recentes → restante
    fila = [r["usuario"] for r in engajados] + [r["usuario"] for r in seguidores]
    fila = [u for u in dict.fromkeys(fila) if u not in ja]
    if a.limite:
        fila = fila[: a.limite]
    if not fila:
        print(f"Nada novo para coletar ({len(ja)} perfis já em perfis.jsonl).")
        return
    if not a.confirmar:
        print(f"Vai coletar {len(fila)} perfis públicos no Apify (lotes de {a.lote}); {len(ja)} já estão salvos.\n"
              f"Isso gasta créditos. Teste com --limite 50 e depois rode com --confirmar.")
        return
    arq = pasta / "perfis.jsonl"
    with arq.open("a", encoding="utf-8") as f:
        for i in range(0, len(fila), a.lote):
            lote = fila[i: i + a.lote]
            try:
                itens = apify({"directUrls": [f"https://www.instagram.com/{u}/" for u in lote],
                               "resultsType": "details", "resultsLimit": 1})
            except urllib.error.HTTPError as e:
                if e.code in (401, 402, 403):
                    sys.exit(f"Apify recusou (HTTP {e.code}): token inválido ou sem créditos. Nada foi perdido; rode de novo depois.")
                print(f"Lote {i // a.lote + 1} falhou (HTTP {e.code}); rode de novo depois.", file=sys.stderr)
                continue
            except Exception as e:  # segue para o próximo lote; os que faltaram entram na próxima execução
                print(f"Lote {i // a.lote + 1} falhou ({e.__class__.__name__}); rode de novo depois.", file=sys.stderr)
                continue
            achados = set()
            for d in itens:
                u = (d.get("username") or "").lower()
                if not u:
                    continue
                achados.add(u)
                links = [x.get("url") for x in (d.get("externalUrls") or []) if isinstance(x, dict) and x.get("url")]
                reg = {
                    "usuario": u, "nome": d.get("fullName") or "", "bio": d.get("biography") or "",
                    "link": d.get("externalUrl") or (links[0] if links else ""), "links": links,
                    "categoria": d.get("businessCategoryName") or "", "conta_comercial": bool(d.get("isBusinessAccount")),
                    "seguidores": d.get("followersCount"), "seguindo": d.get("followsCount"), "posts": d.get("postsCount"),
                    "privado": bool(d.get("private")), "verificado": bool(d.get("verified")),
                    "coletado_em": datetime.now().isoformat(timespec="seconds"),
                }
                f.write(json.dumps(reg, ensure_ascii=False) + "\n")
            for u in lote:
                if u not in achados:  # perfil apagado, renomeado ou bloqueado
                    f.write(json.dumps({"usuario": u, "nao_encontrado": True,
                                        "coletado_em": datetime.now().isoformat(timespec="seconds")}) + "\n")
            f.flush()
            print(f"{min(i + a.lote, len(fila))}/{len(fila)} perfis", file=sys.stderr)
    print(f"{len(carregar_perfis(pasta))} perfis em {arq}")


# ---------- 4. pontuar ----------
def pontuar_perfil(p: dict, icp: dict, seguiu_em: str, comentarios: int, hoje: datetime) -> dict:
    texto = norm(" ".join(p.get(k) or "" for k in ("nome", "bio", "categoria")))
    links = " ".join([p.get("link") or ""] + [x for x in (p.get("links") or []) if x]).lower()
    pontos, motivos = 0, []

    dec = achar(texto, icp["decisor"]["palavras"])
    if dec:
        pontos += icp["decisor"]["peso"]; motivos.append(f"decisor ({', '.join(dec[:3])})")

    emp = icp["empresa"]
    fraca = norm(p.get("categoria")) in {norm(c) for c in emp["categorias_fracas"]}
    if p.get("conta_comercial"):
        pontos += emp["peso_conta_comercial"]; motivos.append("conta comercial")
    if p.get("categoria") and not fraca:
        pontos += emp["peso_categoria"]; motivos.append(f"categoria: {p['categoria']}")
    sinais = achar(texto, emp["palavras"])
    if sinais:
        pontos += emp["peso_palavra"]; motivos.append(f"empresa ({', '.join(sinais[:3])})")

    seg = p.get("seguidores") or 0
    for fx in icp["porte"]["faixas"]:
        if fx["min"] <= seg <= fx["max"]:
            pontos += fx["pontos"]; break

    lk = icp["link"]
    if any(x in links for x in lk["whatsapp"]["palavras"]):
        pontos += lk["whatsapp"]["pontos"]; motivos.append("WhatsApp na bio")
    elif any(x in links for x in lk["agregador"]["palavras"]):
        pontos += lk["agregador"]["pontos"]
    elif links.strip():
        pontos += lk["site"]["pontos"]; motivos.append("site próprio")

    if seguiu_em:
        dias = (hoje - datetime.fromisoformat(seguiu_em).replace(tzinfo=timezone.utc)).days
        if dias <= icp["seguiu_recente"]["dias"]:
            pontos += icp["seguiu_recente"]["pontos"]; motivos.append(f"seguiu há {dias} dias")
    if comentarios:
        pontos += icp["engajou"]["pontos"]; motivos.append(f"comentou {comentarios}x")

    pil_txt = texto + " " + links
    pilares = sorted(((len(achar(pil_txt, kws)), nome) for nome, kws in icp["pilares"].items()), reverse=True)
    pilares = [n for q, n in pilares if q]

    conc = achar(texto, icp["concorrente_ou_parceiro"]["palavras"])
    desc = achar(texto, icp["descartar"]["palavras"])
    seguindo = p.get("seguindo") or 0
    if not desc and seg and seguindo > icp["descartar"]["seguindo_vs_seguidores"] * seg and seg < 300:
        desc = ["segue muito mais do que é seguido"]
    if not desc and (p.get("posts") or 0) < icp["descartar"]["min_posts"]:
        desc = ["sem posts"]

    fx = icp["faixas_de_prioridade"]
    if desc:
        grupo = "descartado"
    elif conc:
        grupo = "parceiro/concorrente"
    elif pontos >= fx["A"]:
        grupo = "A"
    elif pontos >= fx["B"]:
        grupo = "B"
    elif pontos >= fx["C"]:
        grupo = "C"
    else:
        grupo = "fora do perfil"
    return {"pontos": pontos, "grupo": grupo, "motivos": "; ".join(motivos),
            "pilares": ", ".join(pilares) or "a definir",
            "alerta": ", ".join(desc or conc)}


COLUNAS = ["grupo", "pontos", "usuario", "nome", "categoria", "seguidores", "pilares", "motivos", "alerta",
           "bio", "link", "conta_comercial", "privado", "seguiu_em", "comentarios", "perfil_url"]


def cmd_pontuar(a) -> None:
    pasta = pasta_conta(a.conta)
    icp = json.loads(Path(a.icp).read_text(encoding="utf-8"))
    perfis = carregar_perfis(pasta)
    if not perfis:
        sys.exit("Rode antes a etapa 'enriquecer'.")
    seguiu = {r["usuario"]: r["seguiu_em"] for r in ler_csv(pasta / "seguidores.csv")}
    coment = {r["usuario"]: int(r["comentarios"]) for r in ler_csv(pasta / "engajados.csv")}
    hoje = datetime.now(timezone.utc)

    linhas, sem_dados = [], []
    for u, p in perfis.items():
        base = {"usuario": u, "perfil_url": f"https://www.instagram.com/{u}/", "seguiu_em": seguiu.get(u, ""),
                "comentarios": coment.get(u, 0)}
        if p.get("nao_encontrado"):
            sem_dados.append({**base, "grupo": "não encontrado", "pontos": 0})
            continue
        if str(p.get("categoria")).lower() in ("none", "null"):
            p = {**p, "categoria": ""}
        r = pontuar_perfil(p, icp, seguiu.get(u, ""), coment.get(u, 0), hoje)
        linhas.append({**base, **r, "nome": p.get("nome"), "categoria": p.get("categoria"),
                       "seguidores": p.get("seguidores"), "bio": (p.get("bio") or "").replace("\n", " / "),
                       "link": p.get("link"), "conta_comercial": "sim" if p.get("conta_comercial") else "",
                       "privado": "sim" if p.get("privado") else ""})
    ordem = {"A": 0, "B": 1, "C": 2, "parceiro/concorrente": 3, "fora do perfil": 4, "descartado": 5, "não encontrado": 6}
    linhas.sort(key=lambda r: (ordem[r["grupo"]], -r["pontos"]))
    todos = linhas + sem_dados

    gravar_csv(pasta / "leads.csv", todos, COLUNAS)
    gravar_xlsx(pasta / "leads.xlsx", todos)
    (pasta / "resumo.md").write_text(resumo(a.conta, todos, len(seguiu), len(coment)), encoding="utf-8")
    cont = {g: sum(1 for r in todos if r["grupo"] == g) for g in ordem}
    print(" · ".join(f"{g}: {n}" for g, n in cont.items() if n))
    print(f"→ {pasta / 'leads.xlsx'} e resumo.md")


def gravar_xlsx(destino: Path, linhas: list[dict]) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        print("openpyxl não instalado; ficou só o CSV.", file=sys.stderr)
        return
    abas = {"Clientes A": ["A"], "Clientes B": ["B"], "Clientes C": ["C"], "Parceiros": ["parceiro/concorrente"],
            "Fora do perfil": ["fora do perfil"], "Descartados": ["descartado", "não encontrado"]}
    wb = Workbook()
    wb.remove(wb.active)
    for nome, grupos in abas.items():
        ws = wb.create_sheet(nome)
        cols = COLUNAS + ["status", "observações"]
        ws.append(cols)
        for c in ws[1]:
            c.font = Font(bold=True, color="F5F0EB")
            c.fill = PatternFill("solid", fgColor="2E0F36")
        for r in linhas:
            if r["grupo"] in grupos:
                ws.append([r.get(c, "") for c in COLUNAS] + ["", ""])
        for col, larg in zip("ABCDEFGHIJKLMNOPQR", [8, 7, 22, 24, 22, 11, 26, 50, 22, 70, 30, 9, 8, 11, 11, 36, 14, 30]):
            ws.column_dimensions[col].width = larg
        ws.freeze_panes = "D2"
        ws.auto_filter.ref = ws.dimensions
    wb.save(destino)


def resumo(conta: str, linhas: list[dict], n_seg: int, n_eng: int) -> str:
    def n(g):
        return sum(1 for r in linhas if r["grupo"] == g)
    alvos = [r for r in linhas if r["grupo"] in ("A", "B")]
    pil = {}
    for r in alvos:
        for p in r["pilares"].split(", "):
            pil[p] = pil.get(p, 0) + 1
    cat = {}
    for r in alvos:
        if r.get("categoria"):
            cat[r["categoria"]] = cat.get(r["categoria"], 0) + 1
    out = [f"# Prospecção na base de @{conta.lstrip('@')}", "",
           f"Gerado em {datetime.now():%d/%m/%Y %H:%M}. Seguidores na exportação: {n_seg}. "
           f"Perfis analisados: {len(linhas)}. Comentaram nos posts: {n_eng}.", "",
           "| Grupo | Perfis |", "|---|---|"]
    out += [f"| {g} | {n(g)} |" for g in ["A", "B", "C", "parceiro/concorrente", "fora do perfil", "descartado", "não encontrado"]]
    out += ["", "## Dor provável (A e B)", "| Pilar | Perfis |", "|---|---|"]
    out += [f"| {k} | {v} |" for k, v in sorted(pil.items(), key=lambda kv: -kv[1])]
    out += ["", "## Categorias mais comuns (A e B)", "| Categoria | Perfis |", "|---|---|"]
    out += [f"| {k} | {v} |" for k, v in sorted(cat.items(), key=lambda kv: -kv[1])[:15]]
    out += ["", "## Top 30", "| # | Perfil | Grupo | Pontos | Pilares | Por quê |", "|---|---|---|---|---|---|"]
    for i, r in enumerate(alvos[:30], 1):
        out.append(f"| {i} | [@{r['usuario']}]({r['perfil_url']}) | {r['grupo']} | {r['pontos']} | {r['pilares']} | {r['motivos']} |")
    out += ["", "> Pontuação automática por palavras da bio, categoria, link e porte (regras em `prospeccao/icp.json`). "
            "É uma triagem: confirme cada perfil antes de abordar."]
    return "\n".join(out) + "\n"


def main() -> None:
    ap = argparse.ArgumentParser(description="Prospecção de clientes na base de seguidores do Instagram.")
    ap.add_argument("--conta", default="odilon.mentor", help="sua conta (padrão: odilon.mentor)")
    sub = ap.add_subparsers(dest="etapa", required=True)

    s = sub.add_parser("seguidores", help="lê a exportação oficial do Instagram (.zip, pasta, .json, .html ou .csv)")
    s.add_argument("arquivo")
    s.set_defaults(f=cmd_seguidores)

    s = sub.add_parser("engajados", help="(Apify) quem comentou nos seus últimos posts")
    s.add_argument("--posts", type=int, default=30)
    s.add_argument("--por-post", type=int, default=200)
    s.add_argument("--confirmar", action="store_true")
    s.set_defaults(f=cmd_engajados)

    s = sub.add_parser("enriquecer", help="(Apify) dados públicos de cada perfil; retoma de onde parou")
    s.add_argument("--limite", type=int, default=0, help="máximo de perfis nesta execução (0 = todos)")
    s.add_argument("--lote", type=int, default=50)
    s.add_argument("--confirmar", action="store_true")
    s.set_defaults(f=cmd_enriquecer)

    s = sub.add_parser("pontuar", help="aplica o ICP e gera leads.xlsx + resumo.md")
    s.add_argument("--icp", default=str(ICP_PADRAO))
    s.set_defaults(f=cmd_pontuar)

    a = ap.parse_args()
    a.f(a)


if __name__ == "__main__":
    main()
