#!/usr/bin/env python3
"""Varredura de perfil do Instagram: dados do perfil, posts, mídias e análise de criativos.

Uso:
    python scraper/instagram_scraper.py fabianocarvalhojr
    python scraper/instagram_scraper.py https://www.instagram.com/fabianocarvalhojr/ --max-posts 100
    python scraper/instagram_scraper.py fabianocarvalhojr --login SEU_USUARIO   # mais posts, menos bloqueio
    python scraper/instagram_scraper.py fabianocarvalhojr --sem-midias
    IG_SESSIONID=... python scraper/instagram_scraper.py fabianocarvalhojr  # usa o cookie sessionid de uma conta logada

Saída em output/<usuario>/:
    perfil.json        dados do perfil
    posts.csv          um post por linha
    <usuario>.xlsx     planilha com abas Perfil, Posts, Top Posts e Resumo
    analise.md         relatório de criativos (ganchos, CTAs, formatos, hashtags)
    midias/            imagens e vídeos baixados
"""

import argparse
import csv
import json
import os
import re
import sys
import time
from collections import Counter
from datetime import datetime
from pathlib import Path

import instaloader
from openpyxl import Workbook
from openpyxl.styles import Font

CTA_PATTERNS = {
    "link na bio": r"link na bio|link da bio|clique no link",
    "comente": r"\bcomente\b|comenta\b|deixa? (aqui )?nos coment",
    "salve": r"\bsalv[ea]\b|salva esse|salve esse",
    "compartilhe": r"compartilh[ea]|manda pra|envia pra|marque (um|alguém|quem)",
    "siga": r"\bsiga\b|segue o perfil|me segue",
    "direct/dm": r"\bdirect\b|\bdm\b|chama no",
    "inscreva-se/garanta": r"inscreva|garanta|vagas|compre|acesse",
}

TYPE_LABELS = {"GraphImage": "foto", "GraphVideo": "reels/vídeo", "GraphSidecar": "carrossel"}


def parse_username(value: str) -> str:
    match = re.search(r"instagram\.com/([^/?#]+)", value)
    return (match.group(1) if match else value).lstrip("@").strip("/")


def hook_of(caption: str) -> str:
    """Primeira linha não vazia da legenda: o gancho escrito."""
    for line in caption.splitlines():
        if line.strip():
            return line.strip()[:200]
    return ""


def ctas_of(caption: str) -> list[str]:
    text = caption.lower()
    return [name for name, pattern in CTA_PATTERNS.items() if re.search(pattern, text)]


def collect(username: str, max_posts: int, download_media: bool, login: str | None, out_dir: Path):
    loader = instaloader.Instaloader(
        dirname_pattern=str(out_dir / "midias"),
        filename_pattern="{date_utc:%Y-%m-%d_%H-%M}_{shortcode}",
        download_video_thumbnails=False,
        download_geotags=False,
        download_comments=False,
        save_metadata=False,
        compress_json=False,
        quiet=True,
    )
    sessionid = os.environ.get("IG_SESSIONID")
    if sessionid:
        loader.context._session.cookies.set("sessionid", sessionid, domain=".instagram.com")
        loader.context.username = loader.test_login()
        if not loader.context.username:
            sys.exit("IG_SESSIONID inválido ou expirado.")
    elif login:
        try:
            loader.load_session_from_file(login)
        except FileNotFoundError:
            loader.interactive_login(login)
            loader.save_session_to_file()

    profile = instaloader.Profile.from_username(loader.context, username)
    profile_data = {
        "usuario": profile.username,
        "nome": profile.full_name,
        "bio": profile.biography,
        "link_bio": profile.external_url,
        "seguidores": profile.followers,
        "seguindo": profile.followees,
        "total_posts": profile.mediacount,
        "verificado": profile.is_verified,
        "conta_comercial": profile.is_business_account,
        "categoria": profile.business_category_name,
        "privado": profile.is_private,
        "coletado_em": datetime.now().isoformat(timespec="seconds"),
    }
    if profile.is_private and not login:
        print("Perfil privado: só dá para coletar os dados do perfil.", file=sys.stderr)
        return profile_data, []

    posts = []
    for i, post in enumerate(profile.get_posts()):
        if i >= max_posts:
            break
        caption = post.caption or ""
        likes, comments = post.likes, post.comments
        posts.append({
            "data": post.date_local.strftime("%Y-%m-%d %H:%M"),
            "url": f"https://www.instagram.com/p/{post.shortcode}/",
            "tipo": TYPE_LABELS.get(post.typename, post.typename),
            "curtidas": likes,
            "comentarios": comments,
            "visualizacoes": post.video_view_count if post.is_video else None,
            "engajamento_%": round((likes + comments) / profile.followers * 100, 3) if profile.followers else None,
            "gancho": hook_of(caption),
            "ctas": ", ".join(ctas_of(caption)),
            "hashtags": " ".join(f"#{h}" for h in post.caption_hashtags),
            "mencoes": " ".join(f"@{m}" for m in post.caption_mentions),
            "tamanho_legenda": len(caption),
            "legenda": caption,
        })
        if download_media:
            loader.download_post(post, target=profile.username)
        time.sleep(1.5)  # reduz a chance de bloqueio por excesso de requisições
        print(f"  {i + 1} posts coletados", end="\r", file=sys.stderr)
    print(file=sys.stderr)
    return profile_data, posts


def write_outputs(profile_data: dict, posts: list[dict], out_dir: Path) -> None:
    out_dir.mkdir(parents=True, exist_ok=True)
    (out_dir / "perfil.json").write_text(json.dumps(profile_data, ensure_ascii=False, indent=2), encoding="utf-8")

    columns = list(posts[0].keys()) if posts else []
    with open(out_dir / "posts.csv", "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=columns)
        writer.writeheader()
        writer.writerows(posts)

    top = sorted(posts, key=lambda p: p["curtidas"] + p["comentarios"], reverse=True)[:10]
    summary = summarize(posts)

    wb = Workbook()
    ws = wb.active
    ws.title = "Perfil"
    for key, value in profile_data.items():
        ws.append([key, value])
    for title, rows in (("Posts", posts), ("Top Posts", top)):
        sheet = wb.create_sheet(title)
        sheet.append(columns)
        for row in rows:
            sheet.append([row[c] for c in columns])
        for cell in sheet[1]:
            cell.font = Font(bold=True)
    sheet = wb.create_sheet("Resumo")
    for key, value in summary.items():
        sheet.append([key, value if isinstance(value, (int, float, str)) else json.dumps(value, ensure_ascii=False)])
    wb.save(out_dir / f"{profile_data['usuario']}.xlsx")

    (out_dir / "analise.md").write_text(render_report(profile_data, posts, top, summary), encoding="utf-8")


def summarize(posts: list[dict]) -> dict:
    if not posts:
        return {}
    by_type: dict[str, list[dict]] = {}
    for p in posts:
        by_type.setdefault(p["tipo"], []).append(p)
    return {
        "posts_analisados": len(posts),
        "media_curtidas": round(sum(p["curtidas"] for p in posts) / len(posts)),
        "media_comentarios": round(sum(p["comentarios"] for p in posts) / len(posts)),
        "engajamento_medio_%": round(sum(p["engajamento_%"] or 0 for p in posts) / len(posts), 3),
        "desempenho_por_formato": {
            t: {"posts": len(ps), "media_curtidas": round(sum(p["curtidas"] for p in ps) / len(ps))}
            for t, ps in by_type.items()
        },
        "ctas_mais_usados": Counter(c for p in posts for c in p["ctas"].split(", ") if c).most_common(),
        "hashtags_mais_usadas": Counter(h for p in posts for h in p["hashtags"].split()).most_common(15),
        "posts_por_dia_da_semana": Counter(
            datetime.strptime(p["data"], "%Y-%m-%d %H:%M").strftime("%A") for p in posts
        ).most_common(),
    }


def render_report(profile: dict, posts: list[dict], top: list[dict], summary: dict) -> str:
    lines = [
        f"# Análise de criativos: @{profile['usuario']}",
        "",
        f"**{profile['nome']}**: {profile['seguidores']:,} seguidores · {profile['total_posts']:,} posts",
        f"Bio: {profile['bio']}",
        f"Link da bio: {profile['link_bio']}",
        "",
    ]
    if not posts:
        return "\n".join(lines + ["Nenhum post coletado."])
    lines += [
        "## Números gerais",
        f"- Posts analisados: {summary['posts_analisados']}",
        f"- Média de curtidas: {summary['media_curtidas']:,} · média de comentários: {summary['media_comentarios']:,}",
        f"- Engajamento médio: {summary['engajamento_medio_%']}%",
        "",
        "## Desempenho por formato",
        *[f"- {t}: {d['posts']} posts, média de {d['media_curtidas']:,} curtidas"
          for t, d in summary["desempenho_por_formato"].items()],
        "",
        "## CTAs mais usados",
        *([f"- {c}: {n}x" for c, n in summary["ctas_mais_usados"]] or ["- nenhum CTA padrão detectado"]),
        "",
        "## Top 10 posts (ganchos que mais performaram)",
    ]
    for i, p in enumerate(top, 1):
        lines += [
            f"### {i}. {p['tipo']} · {p['data']} · {p['curtidas']:,} curtidas · {p['comentarios']:,} comentários",
            f"- Gancho: \"{p['gancho']}\"",
            f"- CTAs: {p['ctas'] or 'nenhum'}",
            f"- {p['url']}",
            "",
        ]
    lines += ["## Hashtags mais usadas", " ".join(h for h, _ in summary["hashtags_mais_usadas"]) or "nenhuma"]
    return "\n".join(lines) + "\n"


def main() -> None:
    parser = argparse.ArgumentParser(description="Varre um perfil do Instagram.")
    parser.add_argument("perfil", help="@usuario ou URL do perfil")
    parser.add_argument("--max-posts", type=int, default=50)
    parser.add_argument("--sem-midias", action="store_true", help="não baixar imagens/vídeos")
    parser.add_argument("--login", help="seu usuário do Instagram (pede a senha na 1ª vez e salva a sessão)")
    parser.add_argument("--saida", default="output", help="pasta de saída (padrão: output/)")
    args = parser.parse_args()

    username = parse_username(args.perfil)
    out_dir = Path(args.saida) / username
    print(f"Varrendo @{username}...", file=sys.stderr)
    try:
        profile_data, posts = collect(username, args.max_posts, not args.sem_midias, args.login, out_dir)
    except instaloader.exceptions.ConnectionException as e:
        sys.exit(f"Instagram recusou a conexão ({e}). Tente de novo mais tarde ou use --login.")
    write_outputs(profile_data, posts, out_dir)
    print(f"Pronto: {len(posts)} posts salvos em {out_dir}/", file=sys.stderr)


if __name__ == "__main__":
    main()
